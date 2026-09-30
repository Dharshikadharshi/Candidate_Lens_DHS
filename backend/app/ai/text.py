"""Resume text extraction, PII minimisation and evidence-quote verification."""
import hashlib
import io
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Optional

MAX_RESUME_CHARS = 24000  # keeps prompts well inside model context limits

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)")
URL_RE = re.compile(r"\b(?:https?://|www\.)\S+", re.IGNORECASE)


class ResumeTextError(Exception):
    pass


@dataclass
class ResumeText:
    text: str  # redacted, page-marked text sent to the model
    pages: int
    chars: int
    truncated: bool
    redactions: dict = field(default_factory=dict)
    github_url: Optional[str] = None
    linkedin_url: Optional[str] = None
    leetcode_url: Optional[str] = None
    hackerrank_url: Optional[str] = None


def file_fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _pdf_pages_and_links(data: bytes) -> tuple[list[str], list[str]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        links = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
            if "/Annots" in page:
                annots = page["/Annots"]
                if annots:
                    for annot in annots:
                        try:
                            obj = annot.get_object()
                            if obj and "/A" in obj and "/URI" in obj["/A"]:
                                links.append(obj["/A"]["/URI"])
                        except Exception:
                            pass
        return pages, links
    except Exception as exc:  # malformed or encrypted PDF
        raise ResumeTextError(f"Could not read the PDF: {exc.__class__.__name__}") from exc


def _docx_pages(data: bytes) -> list[str]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeTextError(f"Could not read the DOCX file: {exc.__class__.__name__}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return ["\n".join(parts)]


def redact(text: str) -> tuple[str, dict]:
    """Remove contact details the assessment does not need before text is sent to the model."""
    counts = {}
    for label, pattern in (("email", EMAIL_RE), ("url", URL_RE)):
        text, counts[label] = pattern.subn(f"[{label} removed]", text)

    phones = 0

    def _phone(match: re.Match) -> str:
        nonlocal phones
        # Require 10+ digits so year ranges such as "2019 - 2023" are kept.
        if sum(ch.isdigit() for ch in match.group(0)) >= 10:
            phones += 1
            return "[phone removed]"
        return match.group(0)

    text = PHONE_RE.sub(_phone, text)
    counts["phone"] = phones
    return text, counts


def extract_github_url(text: str, annotations: list[str] = None) -> Optional[str]:
    github_pattern = re.compile(r"(?:https?://)?(?:www\.)?github\.com/([a-zA-Z0-9-]+)/?", re.IGNORECASE)
    
    if annotations:
        for ann in annotations:
            match = github_pattern.search(ann)
            if match:
                return f"https://github.com/{match.group(1)}"
                
    match = github_pattern.search(text)
    if match:
         return f"https://github.com/{match.group(1)}"
    
    return None

def extract_linkedin_url(text: str, annotations: list[str] = None) -> Optional[str]:
    linkedin_pattern = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/([a-zA-Z0-9_-]+)/?", re.IGNORECASE)
    if annotations:
        for ann in annotations:
            match = linkedin_pattern.search(ann)
            if match:
                return f"https://www.linkedin.com/in/{match.group(1)}"
    match = linkedin_pattern.search(text)
    if match:
        return f"https://www.linkedin.com/in/{match.group(1)}"
    return None

def extract_leetcode_url(text: str, annotations: list[str] = None) -> Optional[str]:
    leetcode_pattern = re.compile(r"(?:https?://)?(?:www\.)?leetcode\.com/(?:u/)?([a-zA-Z0-9_-]+)/?", re.IGNORECASE)
    if annotations:
        for ann in annotations:
            match = leetcode_pattern.search(ann)
            if match:
                return f"https://leetcode.com/{match.group(1)}/"
    match = leetcode_pattern.search(text)
    if match:
        return f"https://leetcode.com/{match.group(1)}/"
    return None

def extract_hackerrank_url(text: str, annotations: list[str] = None) -> Optional[str]:
    hackerrank_pattern = re.compile(r"(?:https?://)?(?:www\.)?hackerrank\.com/([a-zA-Z0-9_-]+)/?", re.IGNORECASE)
    if annotations:
        for ann in annotations:
            match = hackerrank_pattern.search(ann)
            if match:
                return f"https://www.hackerrank.com/{match.group(1)}"
    match = hackerrank_pattern.search(text)
    if match:
        return f"https://www.hackerrank.com/{match.group(1)}"
    return None


def extract_resume_text(data: bytes, file_type: str, filename: str) -> ResumeText:
    name = filename.lower()
    links = []
    if file_type == "application/pdf" or name.endswith(".pdf"):
        pages, links = _pdf_pages_and_links(data)
    elif name.endswith(".docx"):
        pages = _docx_pages(data)
    else:
        raise ResumeTextError("Legacy .doc files cannot be analysed. Please upload the resume as PDF or DOCX.")

    marked = []
    raw_text_for_extraction = ""
    for i, page in enumerate(pages, start=1):
        cleaned = re.sub(r"[ \t]+", " ", page).strip()
        if cleaned:
            marked.append(f"=== Page {i} ===\n{cleaned}" if len(pages) > 1 else cleaned)
            raw_text_for_extraction += cleaned + " "
            
    github_url = extract_github_url(raw_text_for_extraction, links)
    linkedin_url = extract_linkedin_url(raw_text_for_extraction, links)
    leetcode_url = extract_leetcode_url(raw_text_for_extraction, links)
    hackerrank_url = extract_hackerrank_url(raw_text_for_extraction, links)
    
    text = "\n\n".join(marked)
    if len(re.sub(r"\W", "", text)) < 40:
        raise ResumeTextError("No readable text was found in the resume (it may be a scanned image).")

    text, counts = redact(text)
    truncated = len(text) > MAX_RESUME_CHARS
    if truncated:
        text = text[:MAX_RESUME_CHARS]
    return ResumeText(
        text=text, pages=len(pages), chars=len(text), truncated=truncated, redactions=counts, 
        github_url=github_url, linkedin_url=linkedin_url, leetcode_url=leetcode_url, hackerrank_url=hackerrank_url
    )


def normalize(text: str) -> str:
    text = text.lower().replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"[^a-z0-9']+", " ", text).strip()


def quote_in_source(quote: Optional[str], source: str, threshold: float = 0.9) -> bool:
    """True when `quote` appears (near-)verbatim in `source`. Guards against fabricated evidence."""
    if not quote:
        return False
    q, s = normalize(quote), normalize(source)
    if len(q) < 3:
        return False
    if q in s:
        return True
    # Tolerate small differences (punctuation/transcription) with a sliding fuzzy match.
    window = len(q)
    if window > len(s):
        return SequenceMatcher(None, q, s).ratio() >= threshold
    step = max(1, window // 4)
    best = 0.0
    for start in range(0, len(s) - window + 1, step):
        ratio = SequenceMatcher(None, q, s[start:start + window]).ratio()
        best = max(best, ratio)
        if best >= threshold:
            return True
    return False


def similarity(a: str, b: str) -> float:
    ta, tb = set(normalize(a).split()), set(normalize(b).split())
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)
